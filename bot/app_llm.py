from flask import Flask, render_template, request, jsonify
import os
import sys
import json
from werkzeug.utils import secure_filename
from openpyxl import load_workbook

sys.path.append(os.path.join(os.path.dirname(__file__), 'src'))

from llm_chatbot import LLMChatBot, create_bot_from_config

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 8 * 1024 * 1024

try:
    bot = create_bot_from_config()
    print("LLM问答机器人启动成功！")
except Exception as e:
    print(f"警告: 无法加载LLM配置: {e}")
    print("请配置 config.json 文件后重启服务")
    bot = None

@app.route('/')
def index():
    return render_template('index_llm.html')

@app.route('/ask', methods=['POST'])
def ask():
    if bot is None:
        return jsonify({'error': '机器人未正确配置，请检查config.json'}), 500
    
    data = request.get_json()
    question = data.get('question', '')
    
    if not question:
        return jsonify({'error': '请输入问题'}), 400
    
    try:
        answer = bot.ask(question)
        return jsonify({'answer': answer})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/clear', methods=['POST'])
def clear_history():
    if bot is None:
        return jsonify({'error': '机器人未正确配置'}), 500
    
    try:
        message = bot.clear_history()
        return jsonify({'message': message})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/memory', methods=['GET'])
def memory():
    if bot is None:
        return jsonify({'error': '机器人未正确配置'}), 500

    try:
        return jsonify(bot.get_memory())
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/memory/clear', methods=['POST'])
def clear_memory():
    if bot is None:
        return jsonify({'error': '机器人未正确配置'}), 500

    try:
        message = bot.clear_memory()
        return jsonify({'message': message, 'memory': bot.get_memory()})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/memory/reset', methods=['POST'])
def reset_memory():
    if bot is None:
        return jsonify({'error': '机器人未正确配置'}), 500

    try:
        message = bot.reset_memory()
        return jsonify({'message': message, 'memory': bot.get_memory()})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/upload', methods=['POST'])
def upload_document():
    if bot is None:
        return jsonify({'error': '机器人未正确配置'}), 500

    uploaded_file = request.files.get('file')
    if uploaded_file is None or not uploaded_file.filename:
        return jsonify({'error': '请选择要上传的 xlsx 文件'}), 400

    filename = secure_filename(uploaded_file.filename)
    if not filename.lower().endswith('.xlsx'):
        return jsonify({'error': '目前只支持 .xlsx 文件'}), 400

    try:
        content, diagnostics = xlsx_to_text(uploaded_file)
        if not content.strip():
            return jsonify({
                'error': 'xlsx 文件没有解析到有效内容，可能是图片版成绩单或表格内容不是普通单元格。',
                'diagnostics': diagnostics,
            }), 400
        if diagnostics.get('non_empty_cells', 0) == 0:
            return jsonify({
                'error': 'xlsx 里没有读到普通文本/数字单元格；如果成绩单是截图嵌入 Excel，需要先转成可复制的表格文本。',
                'diagnostics': diagnostics,
            }), 400
        message = bot.add_document_memory(filename, content)
        return jsonify({
            'message': message,
            'filename': filename,
            'characters': len(content),
            'diagnostics': diagnostics,
            'memory': bot.get_memory(),
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/memory/text', methods=['POST'])
def add_memory_text():
    if bot is None:
        return jsonify({'error': '机器人未正确配置'}), 500

    data = request.get_json() or {}
    title = (data.get('title') or '手动粘贴的学生信息').strip()
    content = (data.get('content') or '').strip()

    if not content:
        return jsonify({'error': '请输入要写入长期记忆的文本'}), 400

    try:
        message = bot.add_document_memory(title, content)
        return jsonify({
            'message': message,
            'title': title,
            'characters': len(content),
            'memory': bot.get_memory(),
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

def xlsx_to_text(file_storage):
    workbook = load_workbook(file_storage, data_only=True, read_only=False)
    content, diagnostics = workbook_to_text(workbook)
    if diagnostics["non_empty_cells"] > 0:
        return content, diagnostics

    file_storage.stream.seek(0)
    formula_workbook = load_workbook(file_storage, data_only=False, read_only=False)
    formula_content, formula_diagnostics = workbook_to_text(formula_workbook)
    formula_diagnostics["data_only_empty"] = True
    return formula_content, formula_diagnostics

def workbook_to_text(workbook):
    parts = []
    diagnostics = {
        "sheets": [],
        "non_empty_cells": 0,
        "image_count": 0,
    }

    for sheet in workbook.worksheets:
        image_count = len(getattr(sheet, "_images", []))
        diagnostics["image_count"] += image_count

        rows = list(sheet.iter_rows(values_only=True))
        sheet_non_empty = sum(1 for row in rows for value in row if cell_to_text(value))
        diagnostics["non_empty_cells"] += sheet_non_empty
        diagnostics["sheets"].append({
            "title": sheet.title,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "non_empty_cells": sheet_non_empty,
            "image_count": image_count,
        })

        parts.append(f"成绩单工作表：{sheet.title}")
        if image_count:
            parts.append(f"提示：该工作表包含 {image_count} 张嵌入图片，图片内容无法直接作为单元格解析。")
        if not rows:
            continue

        headers = [cell_to_text(value) for value in rows[0]]
        has_headers = any(headers)

        for row_index, row in enumerate(rows[1:] if has_headers else rows, start=2 if has_headers else 1):
            values = [cell_to_text(value) for value in row]
            if not any(values):
                continue

            if has_headers:
                fields = []
                for header, value in zip(headers, values):
                    if header and value:
                        fields.append(f"{header}: {value}")
                    elif value:
                        fields.append(value)
                row_text = "；".join(fields)
            else:
                row_text = "；".join(value for value in values if value)

            parts.append(f"第{row_index}行：{row_text}")

    return "\n".join(parts), diagnostics

def cell_to_text(value):
    if value is None:
        return ""
    return str(value).strip()

@app.route('/stats', methods=['GET'])
def stats():
    if bot is None:
        return jsonify({'error': '机器人未正确配置'}), 500
    
    return jsonify({
        'total_triples': len(bot.kg.triples),
        'total_subjects': len(bot.kg.get_all_subjects()),
        'total_relations': len(bot.kg.get_all_relations()),
        'total_objects': len(bot.kg.get_all_objects()),
        'conversation_length': len(bot.conversation_history),
        'memory_items': sum(len(value) for value in bot.get_memory().values() if isinstance(value, list))
    })

@app.route('/graph')
def graph():
    return render_template('graph.html')

@app.route('/graph/data', methods=['GET'])
def graph_data():
    if bot is None:
        return jsonify({'error': '机器人未正确配置'}), 500
    
    # 获取过滤参数
    filter_type = request.args.get('filter', 'all')
    limit = int(request.args.get('limit', 100))
    
    # 根据过滤条件选择三元组
    triples = bot.kg.triples
    if filter_type == 'major':
        triples = [t for t in triples if '专业' in t['subject_type'] or '专业' in t['object']]
    elif filter_type == 'course':
        triples = [t for t in triples if '课程' in t['subject_type'] or '课程' in t['object_type']]
    
    # 限制数量
    triples = triples[:limit]
    
    # 构建节点和边
    nodes = {}
    edges = []
    
    for triple in triples:
        # 添加主体节点
        if triple['subject'] not in nodes:
            nodes[triple['subject']] = {
                'id': triple['subject'],
                'label': triple['subject'],
                'type': triple['subject_type'],
                'group': triple['subject_type']
            }
        
        # 添加客体节点
        if triple['object'] not in nodes:
            nodes[triple['object']] = {
                'id': triple['object'],
                'label': triple['object'],
                'type': triple['object_type'],
                'group': triple['object_type']
            }
        
        # 添加边
        edges.append({
            'from': triple['subject'],
            'to': triple['object'],
            'label': triple['relation'],
            'title': f"{triple['subject']} {triple['relation']} {triple['object']}"
        })
    
    return jsonify({
        'nodes': list(nodes.values()),
        'edges': edges
    })

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)
