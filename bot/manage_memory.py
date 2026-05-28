import argparse
import json
import os
import sys
from datetime import datetime

from openpyxl import load_workbook

sys.path.append(os.path.join(os.path.dirname(__file__), "src"))

from graphiti_memory_store import GraphitiMemoryStore


def load_memory_store():
    config_file = os.path.join(os.path.dirname(__file__), "config.json")
    if not os.path.exists(config_file):
        raise FileNotFoundError("config.json 不存在，请先从 config.example.json 复制并填写配置。")

    with open(config_file, "r", encoding="utf-8") as file:
        config = json.load(file)

    memory_config = config.get("memory", {})
    return GraphitiMemoryStore(memory_config)


def xlsx_to_text(path):
    workbook = load_workbook(path, data_only=True, read_only=False)
    content, diagnostics = workbook_to_text(workbook, path)
    if diagnostics["non_empty_cells"] > 0:
        return content

    workbook = load_workbook(path, data_only=False, read_only=False)
    content, diagnostics = workbook_to_text(workbook, path)
    if diagnostics["non_empty_cells"] == 0:
        raise ValueError(f"xlsx 没有读到普通单元格内容，诊断信息：{diagnostics}")
    return content


def workbook_to_text(workbook, path):
    parts = [f"成绩单文件：{os.path.basename(path)}"]
    diagnostics = {
        "sheets": [],
        "non_empty_cells": 0,
        "image_count": 0,
    }

    for sheet in workbook.worksheets:
        image_count = len(getattr(sheet, "_images", []))
        diagnostics["image_count"] += image_count

        rows = list(sheet.iter_rows(values_only=True))
        sheet_non_empty = sum(1 for row in rows for value in row if cell_to_text(value))
        diagnostics["non_empty_cells"] += sheet_non_empty
        diagnostics["sheets"].append({
            "title": sheet.title,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "non_empty_cells": sheet_non_empty,
            "image_count": image_count,
        })

        parts.append(f"成绩单工作表：{sheet.title}")
        if image_count:
            parts.append(f"提示：该工作表包含 {image_count} 张嵌入图片，图片内容无法直接作为单元格解析。")
        if not rows:
            continue

        headers = [cell_to_text(value) for value in rows[0]]
        has_headers = any(headers)

        data_rows = rows[1:] if has_headers else rows
        start_index = 2 if has_headers else 1
        for offset, row in enumerate(data_rows):
            values = [cell_to_text(value) for value in row]
            if not any(values):
                continue

            if has_headers:
                fields = []
                for header, value in zip(headers, values):
                    if header and value:
                        fields.append(f"{header}: {value}")
                    elif value:
                        fields.append(value)
                row_text = "；".join(fields)
            else:
                row_text = "；".join(value for value in values if value)

            parts.append(f"第{start_index + offset}行：{row_text}")

    return "\n".join(parts), diagnostics


def cell_to_text(value):
    if value is None:
        return ""
    return str(value).strip()


def main():
    parser = argparse.ArgumentParser(description="Manage Graphiti student memory.")
    parser.add_argument("--reset", action="store_true", help="物理清空当前 group_id 的 Graphiti 记忆")
    parser.add_argument("--xlsx", help="上传成绩单 xlsx 到 Graphiti 记忆")
    parser.add_argument("--intro", help="重新介绍你的个人信息并写入 Graphiti 记忆")
    args = parser.parse_args()

    store = load_memory_store()

    if args.reset:
        store.clear(hard=True)
        print("已物理清空当前 Graphiti group 的长期记忆。")

    if args.xlsx:
        content = xlsx_to_text(args.xlsx)
        store.add_document(
            title=os.path.basename(args.xlsx),
            content=content,
            source_description="Uploaded transcript xlsx",
        )
        print(f"已写入成绩单：{args.xlsx}")

    if args.intro:
        store.add_document(
            title=f"学生自我介绍 {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            content=args.intro,
            source_description="Student self introduction",
        )
        print("已写入自我介绍。")

    print(json.dumps(store.get(), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
